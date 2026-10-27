import pandas as pd, numpy as np, os, gzip
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
import sys; sys.path.insert(0,os.path.dirname(os.path.abspath(__file__))); from headers import H
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
SRC=str(ROOT/'results'); OUT=str(ROOT/'tables'/'output'); os.makedirs(OUT,exist_ok=True)
rd=lambda f,s: pd.read_excel(f'{SRC}/{f}.xlsx',sheet_name=s)
STG={'seedling_leaf':'SL','fruiting_leaf':'FL','fruit_3d':'F3','fruit_6d':'F6','fruit_9d':'F9'}
def stg(d,c='stage'):
    if c in d: d[c]=d[c].map(lambda v: STG.get(v,v))
    return d
BOOK={}
def write(name,title,sheets,notes):
    """sheets: list of (sheet_name, desc, df)"""
    p=f'{OUT}/{name}.xlsx'
    info=[[title],[''],['工作表 Sheet','内容 Content']]+[[s,d] for s,d,_ in sheets]+[['']]+[[n] for n in notes]
    miss=[c for _,_,df in sheets for c in df.columns if c not in H]
    assert not miss,(name,miss)
    sheets=[(s,d,df.rename(columns=H)) for s,d,df in sheets]
    with pd.ExcelWriter(p,engine='openpyxl') as w:
        pd.DataFrame(info).to_excel(w,sheet_name='说明',index=False,header=False)
        for s,d,df in sheets: df.to_excel(w,sheet_name=s,index=False)
    wb=load_workbook(p)
    ws=wb['说明']; ws['A1'].font=Font(bold=True,size=12); ws.column_dimensions['A'].width=24; ws.column_dimensions['B'].width=90
    for c in ws[3]: c.font=Font(bold=True)
    for s,_,df in sheets:
        ws=wb[s]; ws.freeze_panes='A2'; ws.row_dimensions[1].height=44
        for c in ws[1]:
            c.font=Font(bold=True); c.alignment=Alignment(wrap_text=True,vertical='top')
            # Pandas versions differ in their default header border styling.
            # Match the supplied publication tables explicitly.
            c.border=Border(left=Side(),right=Side(),top=Side(),bottom=Side(),diagonal=Side())
        for j,col in enumerate(df.columns,1):
            L=max([len(str(col))*0.6]+[len(str(v)) for v in df[col].head(200)]); ws.column_dimensions[ws.cell(1,j).column_letter].width=min(max(10,L*1.1),40)
            low=str(col).lower()
            if low.startswith(('q','p值','p ','p,','bh','检验族内','跨135','lsd检验','中亲对比','高亲对比','welch检验','组合间q','md组合内','ms组合内','rna-seq q')) or 'q（' in low or 'q (' in low:
                for r in range(2,ws.max_row+1):
                    v=ws.cell(r,j).value
                    if isinstance(v,float) and v!=0 and abs(v)<0.001: ws.cell(r,j).number_format='0.00E+00'
    wb.save(p); BOOK[name]=(title,[(s,d,len(df)) for s,d,df in sheets])

# ---------- S1 ----------
m=rd('S1','样品映射'); q=rd('S1','比对质控')
s1=m.merge(q[['sample_id','input_read_pairs','unique_mapped_pct','multimapped_pct','gene_assigned_fragments','gene_assignment_pct']],on='sample_id')
s1=stg(s1)[['sample_id','library_id','genotype','stage','replicate','input_read_pairs','unique_mapped_pct','multimapped_pct','gene_assigned_fragments','gene_assignment_pct']]
code=rd('S1','样品编码说明')
write('S1','表S1 75个RNA-seq文库的样品信息与比对统计  Table S1 Sample information and alignment statistics of the 75 RNA-seq libraries',
 [('样品与比对','样品编号、基因型、样品类型、重复及STAR比对和基因分配统计',s1)],
 ['样品类型：SL，苗期叶；FL，结瓜期叶；F3、F6、F9，开花后3、6、9 d果实。','表头为中英文对照。',
  '样品数字编码：'+'；'.join(f'{a}={b}' for a,b in zip(code.iloc[:,0],code.iloc[:,1]))])

# ---------- S2 ----------
blk=rd('S2','区组观测')[['block','genotype','plant_count','total_fruit_count','total_yield_kg','yield_g_per_plant','fruits_per_plant','fruit_weight_g']]
an=rd('S2','ANOVA')[['trait','term','sum_squares','df','mean_square','F','P']]
mean=rd('S2','调整均值')[['trait','genotype','n_blocks','adjusted_mean','SD_across_blocks','model_SE','CI95_mean_lower','CI95_mean_upper']]
lsd=rd('S2','LSD全部比较')[['trait','group1','group2','P_LSD']]
h=rd('S2','杂种优势双侧对比'); pot=rd('S2','显性度区间与检验敏感性')
het=h[['trait','hybrid','F1_mean','MPV','higher_parent','higher_parent_mean','MPH_percent','BPH_percent','F1_minus_MPV','CI95_MPV_difference_lower','CI95_MPV_difference_upper','P_MPV_two_sided','P_MPV_Holm2','F1_minus_HP','CI95_HP_difference_lower','CI95_HP_difference_upper','P_BPH_two_sided','P_BPH_two_sided_Holm2']]
het=het.merge(pot[['trait','hybrid','potence','Fieller_unadjusted_CI95_lower','Fieller_unadjusted_CI95_upper']],on=['trait','hybrid'])
het['potence']=het['potence'].astype(object)
het.loc[het.Fieller_unadjusted_CI95_lower.isna(),'potence']='不列示 NR'
dec=rd('S2','乘积完整分解')[['hybrid','MP_yield_g','F1_yield_g','observed_MPH_difference_g','midparent_product_g','manuscript_baseline_gap_g','manuscript_ratio_percent','exact_parent_complement_term_g','fruit_number_departure_term_g','fruit_mass_departure_term_g','departure_product_term_g','block_covariance_term_g','exact_decomposition_sum_g']]
cum=rd('S2','累计采收')[['block','genotype','harvest_index','plant_count','cumulative_yield_g_per_plant','cumulative_fruits_per_plant']]
fl=rd('S2','果长汇总')[['genotype','n_records','mean_cm','SD_cm','min_cm','max_cm']]
write('S2','表S2 产量构成性状、杂种优势检验与产量中亲差分解  Table S2 Yield components, heterosis tests and decomposition of the yield mid-parent difference',
 [('区组观测','各区组×基因型的株数、总瓜数、总产量及单株产量、单株瓜数和单瓜重（分析单位）',blk),
  ('方差分析','性状～区组＋基因型的方差分析（误差df=8）',an),
  ('基因型均值','5个基因型的模型均值、区组间SD、标准误及95%置信区间',mean),
  ('LSD比较','各性状基因型两两LSD比较的P值',lsd),
  ('杂种优势与显性度','中亲/高亲对比（双侧t检验，每个性状内2个组合Holm校正）及显性度h的Fieller 95%区间（未校正）。区间无界时按正文规则不列示h的点估计（标为“不列示 NR”），区间端点留空',het),
  ('产量中亲差分解','产量中亲差的完整代数分解（g/株）；“B占观察中亲差”即正文88.7%与21.0%',dec),
  ('累计采收','各区组逐次累计单株产量和单株瓜数（图S1数据）',cum),
  ('瓜长','每个基因型20条商品瓜的瓜长描述统计（描述性，无区组信息）',fl)],
 ['MPV，中亲值；HP，高亲值；MPH/BPH，中亲/超亲优势（%）。','分解各项：亲本互补项 −(N̄C−N̄P)(W̄C−W̄P)/4、瓜数偏离项 W_MP·δN、单瓜重偏离项 N_MP·δW、偏离乘积项 δN·δW、协方差校正项 c_F1−(c_C+c_P)/2，之和等于观察中亲差。'])

# ---------- S3 ----------
pat=rd('S3','表达模式汇总')
ev=stg(rd('S3','表达偏离事件'))[['gene_id','genotype','stage','log2FoldChange_midparent','padj_midparent','log2FoldChange_hybrid_C','padj_hybrid_C','log2FoldChange_hybrid_paternal','padj_hybrid_paternal','log2FoldChange_parent','padj_parent','publication_class']]
ev=ev.rename(columns={'log2FoldChange_midparent':'D','padj_midparent':'q_D'})
EC=['set','direction','database','term_id','term_name','background_n','selected_n','term_background_n','overlap_n','fold_enrichment','pvalue','q_BH','gene_ids']
sig=rd('S3','显著功能条目')[EC]
sh=rd('S3','共有基因富集'); sh=sh[sh.q_BH<0.05].copy(); sh['direction']=sh.set.str.replace('shared_',''); sh=sh[[c for c in EC if c in sh.columns]]
enr=pd.concat([sig,sh],ignore_index=True)
ov=rd('S3','叶片重叠汇总')
un=rd('S3','叶片偏离基因并集'); un=un[un.overlap_category.str.startswith('both')]
un=un[['gene_id','overlap_category','log2FoldChange_midparent_MD','padj_midparent_MD','log2FoldChange_midparent_MS','padj_midparent_MS','Swiss-Prot_annotation']].rename(columns={'log2FoldChange_midparent_MD':'D_MD','padj_midparent_MD':'q_MD','log2FoldChange_midparent_MS':'D_MS','padj_midparent_MS':'q_MS'})
g=rd('S3','图3C基因效应')
lab=g.drop_duplicates('gene_id').set_index('gene_id').label
w1=g.pivot(index='gene_id',columns='comparison',values='log2FC'); w2=g.pivot(index='gene_id',columns='comparison',values='padj')
order=['midparent_MD','midparent_MS','hybrid_C_MD','hybrid_paternal_MD','hybrid_C_MS','hybrid_paternal_MS']; nm={'midparent_MD':'D_MD','midparent_MS':'D_MS','hybrid_C_MD':'MD/C','hybrid_paternal_MD':'MD/D','hybrid_C_MS':'MS/C','hybrid_paternal_MS':'MS/S'}
f3=pd.DataFrame({'gene_id':w1.index,'label':lab.reindex(w1.index).values})
for c in order: f3[nm[c]]=w1[c].values; f3['q '+nm[c]]=w2[c].values
f3['set']=np.where(f3.label.str.contains('LHC'),'MD+ antenna proteins',np.where(f3.label=='PAL','shared negative: PAL activity','shared negative: plant–pathogen interaction'))
f3=f3.set_index('gene_id').loc[g.gene_id.drop_duplicates()].reset_index()
f3=f3[['gene_id','label','set']+[c for c in f3.columns if c not in ('gene_id','label','set')]]
write('S3','表S3 结瓜期叶及各样品类型的F₁表达偏离与功能富集  Table S3 F₁ expression departures and functional enrichment',
 [('表达模式汇总','各组合×样品类型的表达偏离事件按5类模式计数',pat),
  ('表达偏离事件','5 179个表达偏离事件（BH校正q<0.05且|D|≥1）的效应、q值及分类',ev),
  ('叶片显著富集条目','MD、MS结瓜期叶正/负偏离集合及两组合共同正/负集合的显著GO/KEGG条目（检验族内BH q<0.05）及贡献基因',enr),
  ('叶片偏离重叠汇总','两组合结瓜期叶偏离集合的重叠计数（共同可检验背景16 938个基因）',ov),
  ('共有偏离基因','两组合结瓜期叶均达到阈值的315个基因及方向',un),
  ('图3C基因','图3C中23个代表基因的表达偏离D和F₁–亲本log₂FC及q值',f3)],
 ['D = log₂(μF₁) − [log₂(μC)+log₂(μP)]/2；P在MD中为D、在MS中为S。log₂FC（父本/C）为父本相对C的表达差异。','样品类型：SL，苗期叶；FL，结瓜期叶；F3、F6、F9，开花后3、6、9 d果实。','全基因水平的DESeq2结果、全部富集检验及背景随分析数据公开。'])

# ---------- S4 ----------
s4a=rd('S4','阶段交互汇总')
s4b=rd('S4','显著交互事件')[['gene_id','hybrid','comparison','log2FoldChange','lfcSE','pvalue','padj']].rename(columns={'log2FoldChange':'ΔD'})
s4c=rd('S4','显著功能条目')[EC]
write('S4','表S4 果实取样时期间表达偏离变化及功能富集  Table S4 Changes in expression departures between fruit stages and functional enrichment',
 [('阶段交互汇总','各组合三个取样时期对比的可检验基因数及显著偏离变化事件数',s4a),
  ('显著交互事件','表达偏离发生显著变化的事件（各对比内BH q<0.05且|ΔD|≥1）',s4b),
  ('显著富集条目','9 d对6 d正/负偏离变化集合的显著GO/KEGG条目（图S6）',s4c)],
 ['ΔD为两个取样时期间F₁相对亲本对数均值偏离的变化；正值表示偏离随取样时期增大。'])

# ---------- S5 (unchanged content) ----------
write('S5','表S5 亲本差异SNP与共同SNP筛选  Table S5 Selection of diagnostic and shared SNPs',
 [('亲本差异SNP筛选','C–D和C–S亲本间纯合差异SNP的逐步筛选数目',rd('S5','亲本诊断SNP筛选')),('共同SNP筛选','共同SNP及唯一外显子归属位点数（亲本支持≥3或≥4个文库）',rd('S5','共同SNP筛选'))],[])

# ---------- S6 ----------
dg=stg(rd('S6','逐基因诊断'))[['profile','model','genotype','stage','testable','called']]
dg['model']=dg.model.map({'legacy_gene_wise_MLE':'gene-wise','shared_dispersion_simulation':'shared'})
jt=stg(rd('S6','共同可检验分母'))
ev35=stg(rd('S6','共享模型35事件'))
t13=pd.read_csv(f'{SRC}/S6/TableS13_parent_support_sensitivity.tsv',sep='\t'); t13=stg(t13)
ev35=ev35.merge(t13[['stage','gene_id','padj_MLE3','differential_ase_MLE3']],on=['stage','gene_id'],how='left')
ev35['both_models']=ev35.differential_ase&ev35.differential_ase_MLE3.fillna(False).astype(bool)
ev35=ev35[['gene_id','stage','MD_total_depth','MS_total_depth','MD_maternal_fraction','MS_maternal_fraction','delta_log2_ratio_MS_minus_MD','pvalue','padj','padj_MLE3','both_models']].rename(columns={'delta_log2_ratio_MS_minus_MD':'ΔR','padj':'q_shared','padj_MLE3':'q_gene_wise'})
reps=['CsaV3_7G007510','CsaV3_6G002590']
sh21=stg(pd.read_csv(f'{SRC}/S6/TableS21_complete_shared_ASE.tsv.gz',sep='\t',usecols=['genotype','stage','gene_id','testable','total_depth','maternal_fraction','padj','ase_significant']))
ml20=stg(pd.read_csv(f'{SRC}/S6/TableS20_complete_MLE_ASE.tsv.gz',sep='\t',usecols=['genotype','stage','gene_id','testable','total_depth','maternal_fraction','padj','ase_significant']))
rows=[]
for gname in reps:
    for st in ['SL','FL','F3','F6','F9']:
        r={'gene_id':gname,'stage':st}
        c=t13[(t13.gene_id==gname)&(t13.stage==st)]
        r['cross_q_shared']=c.padj_shared3.iloc[0] if len(c) else 'NT'; r['cross_q_gene_wise']=c.padj_MLE3.iloc[0] if len(c) else 'NT'
        for hy in ['MD','MS']:
            for tag,tb in [('shared',sh21),('gene_wise',ml20)]:
                d=tb[(tb.gene_id==gname)&(tb.stage==st)&(tb.genotype==hy)]
                ok=len(d) and bool(d.testable.iloc[0])
                r[f'{hy}_C_fraction_{tag}']=d.maternal_fraction.iloc[0] if ok else 'NT'
                r[f'{hy}_q_{tag}']=d.padj.iloc[0] if ok else 'NT'
        rows.append(r)
rep=pd.DataFrame(rows)
fam=rd('S6','40个模拟家族索引')[['family_id','SNP_scope','test_type','genotype','stage','null_hypothesis','null_p0_values','simulations','rejection_count_005','rejection_fraction_005','rejection_ci_low','rejection_ci_high']]; fam=stg(fam)
write('S6','表S6 组合内与组合间ASE检验结果及零假设模拟诊断  Table S6 Within- and cross-hybrid ASE results and null-model diagnostics',
 [('两模型检出数','各SNP范围、模型、组合和样品类型的可检验基因数及组合内ASE检出数（图S3A）。SNP范围：all为各组合全部亲本差异SNP；common3、common4为共同SNP（亲本支持≥3或≥4个文库）',dg),
  ('共同可检验ASE比例','共同SNP（亲本支持≥3个文库）上两组合均可检验基因中的组合内ASE数及比例（图4C）',jt),
  ('组合间35事件','共同SNP（亲本支持≥3个文库）上共享模型检出的35个组合间差异ASE事件及逐基因模型q值；“双模型共同支持”列标明是否两种模型均检出',ev35),
  ('代表基因检验','图5两个代表基因在5类样品中的检验结果。组合间q值基于共同SNP；组合内C等位比例和q值基于各组合全部亲本差异SNP，与图5中共同SNP上的等位比例范围不同，不宜直接比较；NT表示未达可检验条件',rep),
  ('零假设模拟诊断','40组零假设模拟（F01～F40）的SNP范围、设计、名义P<0.05拒绝数及Clopper–Pearson 95%区间（图S3B）',fam)],
 ['C等位比例：母本C来源等位基因的表达比例（以RNA片段计）。','共享模型：跨基因共享离散度的beta-binomial条件模拟；逐基因模型：逐基因离散度的beta-binomial似然比检验。q为BH校正值（组合内按组合×样品类型，组合间按样品类型）。','各工作表的SNP范围见上表“内容”栏。','全部基因的完整ASE结果随分析数据公开。'])

# ---------- S7 ----------
c8=rd('S7','8候选整合证据')
te=rd('S7','ASE同情境总表达')
w=te.pivot_table(index=['gene_id','stage'],columns='genotype',values=['log2FoldChange_parent','log2FoldChange_midparent','padj_midparent'],aggfunc='first')
w.columns=[f'{a}_{b}' for a,b in w.columns]; w=w.reset_index()
w=w.rename(columns={'log2FoldChange_parent_MD':'log2FC_D_vs_C','log2FoldChange_parent_MS':'log2FC_S_vs_C','log2FoldChange_midparent_MD':'D_MD','log2FoldChange_midparent_MS':'D_MS','padj_midparent_MD':'q_D_MD','padj_midparent_MS':'q_D_MS'})
c8=c8.merge(w,on=['gene_id','stage'],how='left'); c8=stg(c8)
c8=c8[['gene_id','stage','MD_maternal_fraction_Shared','MS_maternal_fraction_Shared','delta_log2_ratio_MS_minus_MD_Shared','padj_Shared','padj_MLE','MD_total_depth','MS_total_depth','MD_n_covered_snps','MS_n_covered_snps','support4_shared_called','support4_MLE_called','D_MD','q_D_MD','D_MS','q_D_MS','log2FC_D_vs_C','log2FC_S_vs_C','Swiss-Prot_annotation','Pfam_annotation','NR_annotation']].rename(columns={'delta_log2_ratio_MS_minus_MD_Shared':'ΔR_shared'})
for c_ in ['support4_shared_called','support4_MLE_called']: c8[c_]=c8[c_].map(lambda v: 'NT' if pd.isna(v) else bool(v))
lib=stg(rd('S7','代表基因逐文库片段'))
write('S7','表S7 8个双模型共同支持的组合间ASE候选  Table S7 Eight cross-hybrid ASE candidates supported by both models',
 [('8个候选','等位比例、ΔR、两模型q值、覆盖深度与SNP数、亲本支持≥4时的检出、同样品类型总表达偏离及亲本间log₂FC、同源注释',c8),
  ('代表基因逐文库片段','两个代表基因各F₁文库在共同SNP上的C/父本等位片段数及C等位比例（图5数据）',lib)],
 ['C等位比例：母本C来源等位基因的表达比例（以RNA片段计）。ΔR = log₂(C/S)MS − log₂(C/D)MD；正值表示MS较MD更偏向母本C来源等位基因的表达。“亲本支持≥4时检出”列标明提高亲本支持阈值后是否仍被检出（NT表示不可检验）。','D为同样品类型F₁相对亲本对数均值的表达偏离；log₂FC（D/C）、log₂FC（S/C）为父本相对C的log₂FC。'])

# ---------- S8 ----------
pr=rd('S8','引物序列')[['gene_id','assay','direction','sequence_5_to_3']]
rp=stg(rd('S8','逐重复数据'))[['gene_id','stage','genotype','replicate','batch','target_ct','reference_ct','delta_ct','relative_expression']]
mn=stg(rd('S8','均值和字母'))[['gene_id','stage','genotype','n_biological','delta_ct_mean','delta_ct_sd','relative_expression_mean','relative_expression_sd','letters_Games_Howell_005']]
wl=stg(rd('S8','Welch总体检验'))[['gene_id','stage','F_Welch','df_numerator','df_denominator','p_Welch','q_BH_across_135_panels']]
mn=mn.merge(wl,on=['gene_id','stage'],how='left')
pe=stg(rd('S8','亲本效应匹配')); pe['effect_set']='F1–parent log2FC'
pe=pe.rename(columns={'qPCR_log2FC':'qPCR_effect','log2FoldChange':'RNAseq_effect','padj':'RNAseq_padj'})
md=stg(rd('S8','对数中亲效应匹配')); md['effect_set']='log mid-parent departure'
md=md.rename(columns={'D_log_parental_mean':'qPCR_effect','log2FoldChange':'RNAseq_effect','padj':'RNAseq_padj'})
cols=['effect_set','gene_id','stage','hybrid','comparison','batch','qPCR_effect','RNAseq_effect','RNAseq_padj','included_in_concordance','RNAseq_result_status']
mt=pd.concat([pe[cols],md[cols]],ignore_index=True)
ex=stg(rd('S8','排除原因汇总'))
write('S8','表S8 qRT-PCR引物、检测数据及与RNA-seq的相关性  Table S8 qRT-PCR primers, data and correlation with RNA-seq',
 [('引物序列','27个目标基因及内参Actin2的引物序列（5′→3′）',pr),
  ('逐重复数据','逐生物学重复的Ct、ΔCt及相对表达量（各基因×样品类型内C的几何均值为1）',rp),
  ('组间检验','各基因×样品类型×基因型的均值、SD、Games–Howell字母（P<0.05）及Welch总体检验（含跨135个面板BH q）',mn),
  ('跨平台效应匹配','540项F₁–亲本效应和270项对数中亲偏离的qRT-PCR与RNA-seq配对；“纳入相关分析”列标明是否纳入Spearman相关分析',mt),
  ('排除原因','未纳入相关分析的基因–样品类型及原因（RNA-seq低计数预过滤）',ex)],
 ['样品类型：SL，苗期叶；FL，结瓜期叶；F3、F6、F9，开花后3、6、9 d果实。','相关检验：Spearman ρ（F₁–亲本效应524项，ρ=0.522；中亲偏离262项，ρ=0.451），基因整块置换99 999次，两项Holm校正P均为4×10⁻⁵。'])

from postprocess_tables import postprocess_tables, revise
postprocess_tables(OUT)
# Keep the companion table index consistent with the final display wording.
book_index = {
    k: [
        revise(int(k[1:]), '说明', 'A1', v[0]),
        [(s, revise(int(k[1:]), '说明', f'B{j+4}', d), n)
         for j, (s, d, n) in enumerate(v[1])],
    ]
    for k, v in BOOK.items()
}
import json; json.dump(book_index,open(f'{OUT}/book.json','w'),ensure_ascii=False,indent=1)
for k,v in BOOK.items(): print(k,[(s,n) for s,_,n in v[1]])
