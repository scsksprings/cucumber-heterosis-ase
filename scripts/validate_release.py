#!/usr/bin/env python3
"""Check manuscript denominators, complete ASE exports and optional build output."""
from pathlib import Path
import argparse, hashlib, json, os
import numpy as np
import pandas as pd
from scipy.stats import spearmanr
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]


def compare_workbooks(actual,expected):
    a=load_workbook(actual,read_only=True,data_only=True)
    b=load_workbook(expected,read_only=True,data_only=True)
    try:
        assert a.sheetnames==b.sheetnames,(actual.name,'sheet names')
        cells=0
        for sheet in a.sheetnames:
            ra=list(a[sheet].values);rb=list(b[sheet].values)
            assert len(ra)==len(rb),(actual.name,sheet,'rows',len(ra),len(rb))
            for i,(x,y) in enumerate(zip(ra,rb),1):
                assert len(x)==len(y),(actual.name,sheet,i,'columns')
                for j,(u,v) in enumerate(zip(x,y),1):
                    if isinstance(u,(int,float)) and isinstance(v,(int,float)):
                        assert np.isclose(u,v,rtol=1e-10,atol=1e-12),(actual.name,sheet,i,j,u,v)
                    else:assert u==v,(actual.name,sheet,i,j,u,v)
                    cells+=1
        return cells
    finally:a.close();b.close()


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--built',action='store_true')
    p.add_argument('--submission-tables',type=Path)
    p.add_argument('--report',type=Path,default=ROOT/'work/validation/release.json')
    a=p.parse_args();checks={}
    read=lambda name:pd.read_csv(ROOT/name,sep='\t')
    meta=read('data/rnaseq/sample_metadata.tsv')
    assert len(meta)==75 and meta.sample_id.is_unique
    assert meta.groupby(['genotype','stage']).size().eq(3).all()
    counts=read('data/rnaseq/star_gene_counts.tsv')
    assert counts.shape==(24317,76) and counts.gene_id.is_unique
    assert set(counts.columns[1:])==set(meta.sample_id)
    vals=counts.iloc[:,1:].to_numpy();assert (vals>=0).all() and (vals==vals.astype(int)).all()
    checks['RNAseq']={'samples':75,'genes':24317,'groups':25}
    f1=set(meta.loc[meta.genotype.isin(['MD','MS']),'sample_id'])
    for profile in ['all','common3','common4']:
        paths=list((ROOT/'data/ase'/profile).glob('*.gene_allele_counts.tsv.gz'))
        assert {x.name.split('.')[0] for x in paths}==f1
        for path in paths:
            d=pd.read_csv(path,sep='\t');assert d.gene_id.is_unique
            assert (d[['maternal_fragments','paternal_fragments']]>=0).all().all()
    checks['allele_counts']={'profiles':3,'libraries_per_profile':30}
    sensitivity=read('results/S6/TableS13_parent_support_sensitivity.tsv').set_index(['stage','gene_id'])
    counts_by_model={}
    for support in [3,4]:
        called=[]
        for model,tag in [('genewise','MLE'),('shared','shared')]:
            d=read(f'results/S6/differential_ASE_{model}_support{support}_all.tsv.gz').set_index(['stage','gene_id'])
            assert d.index.is_unique and d.index.isin(sensitivity.index).all()
            np.testing.assert_allclose(d.padj,sensitivity.loc[d.index,f'padj_{tag}{support}'],rtol=1e-12,atol=1e-15,equal_nan=True)
            assert np.array_equal(d.differential_ase.astype(bool),sensitivity.loc[d.index,f'differential_ase_{tag}{support}'].astype(bool))
            called.append(set(d.index[d.differential_ase]))
            counts_by_model[f'{model}_support{support}']={'records':len(d),'called':len(called[-1])}
        assert len(called[0]&called[1])==(8 if support==3 else 7)
    checks['differential_ASE']=counts_by_model
    wells=read('results/S8/qPCR_biological_replicate_values.tsv')
    assert len(wells)==2025 and wells.gene_id.nunique()==27
    assert wells.groupby(['gene_id','stage','genotype']).size().eq(3).all()
    summary=read('results/S8/qPCR_RNAseq_concordance_summary.tsv').set_index('comparison_type')
    pairs={}
    for kind,field,n in [('parental','qPCR_log2FC',524),('log_midparent','D_log_parental_mean',262)]:
        d=read(f'results/S8/qPCR_RNAseq_{kind}_effect_comparison.tsv') if kind=='parental' else read('results/S8/qPCR_RNAseq_log_midparent_comparison.tsv')
        selected=d[d.included_in_concordance];assert len(selected)==n and selected.gene_id.nunique()==27
        rho=spearmanr(selected.log2FoldChange,selected[field]).statistic
        assert np.isclose(rho,summary.loc[field,'spearman_rho_descriptive'],atol=1e-12)
        pairs[kind]={'pairs':n,'rho':float(rho)}
    checks['qPCR']=pairs
    checksum=ROOT/'checksums.sha256'
    if checksum.exists():
        lines=checksum.read_text().splitlines()
        for line in lines:
            digest,name=line.split('  ',1);target=(ROOT/name).resolve()
            assert target.is_relative_to(ROOT) and target.is_file(),name
            assert hashlib.sha256(target.read_bytes()).hexdigest()==digest,('checksum',name)
        checks['checksums']=len(lines)
    if a.built:
        output=Path(os.environ.get('FIG_OUT', ROOT/'figures/output'))
        main=['Fig1_yield_components','Fig2_expression_patterns','Fig3_leaf_comparison',
              'Fig4_SNP_selection_ASE','Fig5_candidate_allele_fractions','Fig6_qPCR_validation']
        for stem in main:
            for ext in ['png','pdf','svg']:
                path=output/f'{stem}.{ext}'
                assert path.is_file() and path.stat().st_size>1000,path
        for ext in ['svg','png']:
            expected={f'FigS4_qPCR_atlas_p{page}.{ext}' for page in range(1,4)}
            actual={p.name for p in (output/'supplementary').glob(f'FigS4_qPCR_atlas_p*.{ext}')}
            assert actual==expected,('atlas pages',ext,actual)
        expected={f'FigS4_{gene}.pdf' for gene in wells.gene_id.unique()}
        actual={p.name for p in (output/'supplementary/panels').glob('FigS4_*.pdf')}
        assert actual==expected,('atlas gene panels',actual^expected)
        for panel in 'ABCDEF':
            assert (output/'panels'/f'Fig6_{panel}.pdf').stat().st_size>1000
        pheno=Path(os.environ.get('PHENO_OUT',output/'phenotype'))
        for suffix,n in [('A',9),('M',15),('H',6),('Mult',2)]:
            assert len(pd.read_csv(pheno/f'c22_{suffix}.csv'))==n
        changes=json.loads((ROOT/'tables/output/formatting_changes.json').read_text())
        assert sum(row['text_cells_changed'] for row in changes)==33
        assert sum(row['format_cells'] for row in changes)==649
        for i in range(1,7):
            matches=list((output/'supplementary').glob(f'FigS{i}_*.pdf'));assert matches,i
        for i in range(1,9):assert (ROOT/f'tables/output/S{i}.xlsx').stat().st_size>1000
        checks['built']={'main_figures':6,'supplementary_figures':6,'condensed_tables':8,'figure6_panels':6,'atlas_editable_SVG_pages':3,'atlas_PNG_pages':3,'atlas_gene_panels':27,'text_cells_revised':33,'richtext_cells':649}
        if a.submission_tables:
            checks['submission_cells_compared']=sum(compare_workbooks(ROOT/f'tables/output/S{i}.xlsx',a.submission_tables/f'S{i}.xlsx') for i in range(1,9))
    report={'status':'passed','checks':checks}
    a.report.parent.mkdir(parents=True,exist_ok=True);a.report.write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))

if __name__=='__main__':main()
