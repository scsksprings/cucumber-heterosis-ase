#!/usr/bin/env python3
"""Compare rebuilt publication tables and figures with an external submission package."""
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from zipfile import ZipFile
import argparse,copy,hashlib,json,os,subprocess,sys,tempfile
import numpy as np
from PIL import Image
from lxml import etree
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'tables'))
from postprocess_tables import sheet_parts
NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def xml_signature(node):
    if node is None:return None
    return (etree.QName(node).localname,tuple(sorted(node.attrib.items())),node.text or '',tuple(xml_signature(c) for c in node))


def table_compare(actual,expected):
    a=load_workbook(actual,read_only=True,data_only=False)
    b=load_workbook(expected,read_only=True,data_only=False)
    result={'file':actual.name,'cells':0,'numbers':0,'booleans':0,'richtext_cells':0}
    def styles(book):
        return [(book._fonts[s.fontId],book._fills[s.fillId],book._borders[s.borderId],
                 book._alignments[s.alignmentId],book._protections[s.protectionId],
                 s.numFmtId if s.numFmtId<164 else book._number_formats[s.numFmtId-164]) for s in book._cell_styles]
    try:
        assert a.sheetnames==b.sheetnames,(actual.name,'sheets')
        sa,sb=styles(a),styles(b)
        for sheet in a.sheetnames:
            x,y=a[sheet],b[sheet]
            assert (x.max_row,x.max_column)==(y.max_row,y.max_column),(actual.name,sheet,'dimensions')
            for ra,rb in zip(x.iter_rows(),y.iter_rows()):
                for ca,cb in zip(ra,rb):
                    u,v=ca.value,cb.value
                    if isinstance(u,bool) or isinstance(v,bool):assert type(u)==type(v) and u==v;result['booleans']+=1
                    elif isinstance(u,(float,int)) and isinstance(v,(float,int)):
                        assert np.isclose(u,v,rtol=1e-10,atol=1e-12),(actual.name,sheet,ca.coordinate,u,v)
                        result['numbers']+=1
                    else:assert u==v,(actual.name,sheet,getattr(ca,'coordinate','blank'),u,v)
                    ida=getattr(ca,'style_id',0);idb=getattr(cb,'style_id',0)
                    assert sa[ida]==sb[idb],(actual.name,sheet,getattr(ca,'coordinate','blank'),'style')
                    result['cells']+=1
    finally:a.close();b.close()
    # Character-level typography is not exposed by openpyxl's read-only values.
    def rich(path):
        entries={}
        with ZipFile(path) as z:
            for sheet,part in sheet_parts(z):
                tree=etree.fromstring(z.read(part))
                for cell in tree.findall('.//s:c',NS):
                    runs=cell.findall('s:is/s:r',NS)
                    if runs:entries[(sheet,cell.get('r'))]=[xml_signature(run) for run in runs]
        return entries
    x,y=rich(actual),rich(expected)
    assert x==y,(actual.name,'richtext',len(x),len(y))
    result['richtext_cells']=len(x);result['passed']=True
    return result


def compare_pdf(job):
    source,actual,out=job;out.mkdir(parents=True,exist_ok=True)
    def render(path,destination):
        destination.mkdir()
        process=subprocess.run(['gs','-q','-dSAFER','-dBATCH','-dNOPAUSE','-sDEVICE=png16m','-r72',
            '-sOutputFile='+str(destination/'page-%02d.png'),str(path)],capture_output=True,text=True)
        if process.returncode:raise RuntimeError(process.stderr[-1500:])
        return sorted(destination.glob('page-*.png'))
    # Isolate each comparison so a shorter revision cannot inherit old PDF pages.
    with tempfile.TemporaryDirectory(prefix='comparison-',dir=out) as scratch:
        scratch=Path(scratch)
        old,new=render(source,scratch/'submitted'),render(actual,scratch/'rebuilt')
        assert len(old)==len(new) and old,(str(source),'PDF page count')
        pages=[]
        for i,(p,q) in enumerate(zip(old,new),1):
            a=np.array(Image.open(p).convert('RGB'));b=np.array(Image.open(q).convert('RGB'))
            same_shape=a.shape==b.shape
            pages.append({'page':i,'same_dimensions':same_shape,'pixel_identical':bool(same_shape and np.array_equal(a,b)),
                'mean_absolute_channel_difference':float(np.abs(a.astype(float)-b).mean()) if same_shape else None})
    return {'file':source.name,'pages':pages,'pixel_identical':all(p['pixel_identical'] for p in pages)}


def compare_png(source,actual):
    with Image.open(source) as a, Image.open(actual) as b:
        x=np.array(a.convert('RGBA'));y=np.array(b.convert('RGBA'))
        same=x.shape==y.shape
        return {'file':source.name,'submitted_dimensions':list(a.size),'rebuilt_dimensions':list(b.size),
                'pixel_identical':bool(same and np.array_equal(x,y))}


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--submission',type=Path,required=True)
    p.add_argument('--report',type=Path,default=ROOT/'work/validation/presentation.json')
    p.add_argument('--compare-pdfs',action='store_true')
    p.add_argument('--compare-pngs',action='store_true',help='Compare full-resolution PNG pixels, including alpha')
    p.add_argument('--require-identical',action='store_true',help='Fail on any requested figure pixel mismatch')
    a=p.parse_args();out=Path(os.environ.get('FIG_OUT',ROOT/'figures/output'))
    rows=[table_compare(ROOT/f'tables/output/S{i}.xlsx',a.submission/f'附表/S{i}.xlsx') for i in range(1,9)]
    pdfs=[]
    if a.compare_pdfs:
        jobs=[]
        for folder,target in [('正文图',out),('附图',out/'supplementary')]:
            for src in sorted((a.submission/folder).glob('*.pdf')):
                actual=target/src.name;assert actual.is_file(),actual
                jobs.append((src,actual,ROOT/'work/validation/pdf_rasters'/src.stem))
        with ThreadPoolExecutor(max_workers=3) as pool:pdfs=list(pool.map(compare_pdf,jobs))
    pngs=[]
    if a.compare_pngs:
        for folder,target in [('正文图',out),('附图',out/'supplementary')]:
            for src in sorted((a.submission/folder).glob('*.png')):
                actual=target/src.name;assert actual.is_file(),actual
                pngs.append(compare_png(src,actual))
        assert pngs,'No source PNG files found'
    if a.compare_pdfs:assert pdfs,'No source PDF files found'
    identical=all(x['pixel_identical'] for x in pdfs+pngs)
    report={'status':'passed','tables':rows,'table_cells':sum(x['cells'] for x in rows),
            'numeric_cells':sum(x['numbers'] for x in rows),'boolean_cells':sum(x['booleans'] for x in rows),
            'richtext_cells':sum(x['richtext_cells'] for x in rows),'pdf_raster_checks':pdfs,
            'all_compared_PDFs_pixel_identical':all(x['pixel_identical'] for x in pdfs) if pdfs else None,
            'png_checks':pngs,'all_compared_PNGs_pixel_identical':all(x['pixel_identical'] for x in pngs) if pngs else None,
            'scope':'Workbook values, effective styles and rich-text runs; optional full-resolution PNG and 72-dpi PDF comparisons. Pixel differences are reported; --require-identical makes them fatal.'}
    if a.require_identical and not identical:report['status']='failed'
    a.report.parent.mkdir(parents=True,exist_ok=True);a.report.write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k not in ['tables','pdf_raster_checks','png_checks']}))
    if report['status']=='failed':raise SystemExit('Figure pixels differ; see the comparison report')

if __name__=='__main__':main()
